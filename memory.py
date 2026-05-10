"""
MAMISHI AI - Memory & RAG System
Uses ChromaDB for vector storage.

Collections:
  1. mamishi_memory     - facts learned from conversations
  2. mamishi_knowledge  - indexed documents from the knowledge folder

CLI examples:
    python memory.py index
    python memory.py status
    python memory.py search "HEMIS Rule 20"
    python memory.py context "What is Rule 20?"
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path


WORKSPACE = Path.home() / "mamishi-ai-workspace"
MEMORY_DIR = WORKSPACE / "memory"
KNOWLEDGE_DIR = WORKSPACE / "knowledge"
CHROMA_DIR = WORKSPACE / "chromadb"
TRACKER_PATH = MEMORY_DIR / "index_tracker.json"

MEMORY_DIR.mkdir(parents=True, exist_ok=True)
KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
CHROMA_DIR.mkdir(parents=True, exist_ok=True)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".csv",
    ".py",
    ".js",
    ".ts",
    ".cs",
    ".sql",
    ".docx",
    ".xlsx",
    ".json",
    ".xml",
}


def get_chroma_client():
    try:
        import chromadb

        return chromadb.PersistentClient(path=str(CHROMA_DIR))
    except ImportError:
        print("ChromaDB not installed. Run: pip install chromadb", file=sys.stderr)
        sys.exit(1)


def get_collections():
    client = get_chroma_client()
    memory_col = client.get_or_create_collection(
        name="mamishi_memory",
        metadata={"description": "Facts and context learned from conversations"},
    )
    knowledge_col = client.get_or_create_collection(
        name="mamishi_knowledge",
        metadata={"description": "Indexed documents from the knowledge folder"},
    )
    return memory_col, knowledge_col


def extract_text_from_file(filepath: Path) -> str:
    ext = filepath.suffix.lower()
    try:
        if ext == ".pdf":
            import fitz

            doc = fitz.open(str(filepath))
            text = "\n\n".join(
                f"[Page {i + 1}]\n{page.get_text().strip()}"
                for i, page in enumerate(doc)
                if page.get_text().strip()
            )
            return text[:100_000]

        if ext == ".docx":
            from docx import Document

            doc = Document(str(filepath))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())[:100_000]

        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        rows.append(row_text)
            return "\n".join(rows)[:100_000]

        if ext == ".csv":
            with open(filepath, "r", encoding="utf-8", errors="replace", newline="") as handle:
                reader = csv.reader(handle)
                return "\n".join(",".join(row) for row in reader)[:100_000]

        with open(filepath, "r", encoding="utf-8", errors="replace") as handle:
            return handle.read()[:100_000]
    except Exception as exc:
        return f"[Error reading {filepath.name}: {exc}]"


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150):
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end]).strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(words):
            break
        start += max(1, chunk_size - overlap)
    return chunks


def file_hash(filepath: Path) -> str:
    digest = hashlib.md5()
    with open(filepath, "rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_tracker() -> dict:
    if TRACKER_PATH.exists():
        with open(TRACKER_PATH, "r", encoding="utf-8") as handle:
            return json.load(handle)
    return {}


def save_tracker(tracker: dict) -> None:
    with open(TRACKER_PATH, "w", encoding="utf-8") as handle:
        json.dump(tracker, handle, indent=2)


def index_folder(folder: Path | None = None, verbose: bool = True):
    folder = folder or KNOWLEDGE_DIR
    _, knowledge_col = get_collections()
    tracker = load_tracker()

    files = [
        path
        for path in Path(folder).rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    added = 0
    updated = 0
    skipped = 0

    for filepath in files:
        key = str(filepath)
        digest = file_hash(filepath)

        if tracker.get(key) == digest:
            skipped += 1
            if verbose:
                print(f"Skipped (unchanged): {filepath.name}")
            continue

        if verbose:
            print(f"Indexing: {filepath.name}")

        text = extract_text_from_file(filepath)
        chunks = chunk_text(text)

        if key in tracker:
            try:
                existing = knowledge_col.get(where={"source": key})
                if existing.get("ids"):
                    knowledge_col.delete(ids=existing["ids"])
            except Exception:
                pass
            updated += 1
        else:
            added += 1

        ids = []
        documents = []
        metadatas = []
        for i, chunk in enumerate(chunks):
            ids.append(f"{digest}_{i}")
            documents.append(chunk)
            metadatas.append(
                {
                    "source": key,
                    "filename": filepath.name,
                    "chunk": i,
                    "total_chunks": len(chunks),
                    "indexed_at": dt.datetime.now().isoformat(),
                }
            )

        if ids:
            knowledge_col.add(ids=ids, documents=documents, metadatas=metadatas)

        tracker[key] = digest

    save_tracker(tracker)
    return {"added": added, "updated": updated, "skipped": skipped, "knowledge_chunks": knowledge_col.count()}


def save_memory(fact: str, category: str = "general", tags: list | None = None) -> bool:
    memory_col, _ = get_collections()
    fact = str(fact or "").strip()
    if not fact:
      return False

    fact_id = hashlib.md5(fact.encode("utf-8")).hexdigest()
    metadata = {
        "category": category,
        "tags": json.dumps(tags or []),
        "created_at": dt.datetime.now().isoformat(),
    }
    try:
        memory_col.upsert(ids=[fact_id], documents=[fact], metadatas=[metadata])
        return True
    except Exception:
        return False


def get_memories(query: str, n_results: int = 5) -> list[str]:
    memory_col, _ = get_collections()
    try:
        if memory_col.count() == 0:
            return []
        results = memory_col.query(query_texts=[query], n_results=min(n_results, memory_col.count()))
        return results.get("documents", [[]])[0]
    except Exception:
        return []


def search_knowledge(query: str, n_results: int = 5, source_filter: str | None = None) -> list[dict]:
    _, knowledge_col = get_collections()
    try:
        if knowledge_col.count() == 0:
            return []
        kwargs = {"query_texts": [query], "n_results": min(n_results, knowledge_col.count())}
        results = knowledge_col.query(**kwargs)
        docs = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        pairs = []
        for doc, meta in zip(docs, metadatas):
            if source_filter and source_filter.lower() not in str(meta.get("filename", "")).lower():
                continue
            pairs.append({"text": doc, "source": meta.get("filename", "?"), "chunk": meta.get("chunk", 0)})
        return pairs
    except Exception:
        return []


def build_context(query: str) -> dict:
    memories = get_memories(query, n_results=4)
    knowledge = search_knowledge(query, n_results=5)

    parts = []
    if memories:
        parts.append("=== MAMISHI AI MEMORY ===")
        parts.append("Facts learned from earlier conversations:")
        parts.extend(f"- {item}" for item in memories)

    if knowledge:
        parts.append("")
        parts.append("=== RELEVANT DOCUMENTS ===")
        parts.append("From the indexed knowledge base:")
        seen_sources = set()
        for item in knowledge:
            source = item["source"]
            if source not in seen_sources:
                parts.append(f"[From: {source}]")
                seen_sources.add(source)
            parts.append(item["text"])

    return {
        "context": "\n".join(part for part in parts if part is not None).strip(),
        "memory_count": len(memories),
        "knowledge_hits": len(knowledge),
        "has_context": bool(parts),
    }


def extract_and_save_memories(conversation_text: str) -> int:
    facts_to_save = []
    for raw_line in str(conversation_text or "").splitlines():
        line = raw_line.strip()
        lower = line.lower()
        if any(token in lower for token in ["auditing", "audit client", "my client"]):
            facts_to_save.append(("audit_context", line[:300]))
        if any(token in lower for token in ["working on", "building", "developing", "my project"]):
            facts_to_save.append(("project", line[:300]))
        if any(token in lower for token in ["i prefer", "i like", "i always", "i use"]):
            facts_to_save.append(("preference", line[:300]))
        if any(token in lower for token in ["sng", "grant thornton", "sadtu", "hemis", "tuteh", "tega", "wesizwe"]):
            facts_to_save.append(("work_context", line[:300]))

    saved = 0
    for category, fact in facts_to_save:
        if len(fact) > 20 and save_memory(fact, category=category):
            saved += 1
    return saved


def get_status() -> dict:
    memory_col, knowledge_col = get_collections()
    tracker = load_tracker()
    return {
        "memory_count": memory_col.count(),
        "knowledge_chunks": knowledge_col.count(),
        "indexed_files": len(tracker),
        "knowledge_folder": str(KNOWLEDGE_DIR),
        "chroma_dir": str(CHROMA_DIR),
    }


def clear_memory():
    client = get_chroma_client()
    try:
        client.delete_collection("mamishi_memory")
    except Exception:
        pass
    client.get_or_create_collection(
        name="mamishi_memory",
        metadata={"description": "Facts and context learned from conversations"},
    )


def clear_knowledge():
    client = get_chroma_client()
    try:
        client.delete_collection("mamishi_knowledge")
    except Exception:
        pass
    client.get_or_create_collection(
        name="mamishi_knowledge",
        metadata={"description": "Indexed documents from the knowledge folder"},
    )
    save_tracker({})


def main():
    parser = argparse.ArgumentParser(description="MAMISHI AI Memory & RAG System")
    parser.add_argument(
        "command",
        choices=["index", "status", "search", "context", "remember", "clear-memory", "clear-knowledge"],
        help="Command to run",
    )
    parser.add_argument("query", nargs="?", help="Query or text input")
    parser.add_argument("--folder", help="Custom folder to index")
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON")
    args = parser.parse_args()

    if args.command == "index":
        folder = Path(args.folder) if args.folder else KNOWLEDGE_DIR
        result = index_folder(folder, verbose=not args.json)
        if args.json:
            print(json.dumps(result))
        else:
            print("\nIndexing complete:")
            print(f"  Added  : {result['added']}")
            print(f"  Updated: {result['updated']}")
            print(f"  Skipped: {result['skipped']}")
            print(f"  Chunks : {result['knowledge_chunks']}")
        return

    if args.command == "status":
        status = get_status()
        if args.json:
            print(json.dumps(status))
        else:
            print("MAMISHI AI Memory Status")
            print(f"  Memory facts    : {status['memory_count']}")
            print(f"  Document chunks : {status['knowledge_chunks']}")
            print(f"  Indexed files   : {status['indexed_files']}")
            print(f"  Knowledge folder: {status['knowledge_folder']}")
            print(f"  ChromaDB path   : {status['chroma_dir']}")
        return

    if args.command == "search":
        if not args.query:
            print("Usage: python memory.py search 'your query'", file=sys.stderr)
            sys.exit(1)
        results = search_knowledge(args.query, n_results=5)
        if args.json:
            print(json.dumps(results))
        elif not results:
            print("No results found. Index a folder first.")
        else:
            for i, item in enumerate(results, 1):
                print(f"\n[{i}] From: {item['source']} (chunk {item['chunk']})")
                print(item["text"][:300] + "...")
        return

    if args.command == "context":
        if not args.query:
            print("Usage: python memory.py context 'your query'", file=sys.stderr)
            sys.exit(1)
        print(json.dumps(build_context(args.query)))
        return

    if args.command == "remember":
        text = args.query if args.query is not None else sys.stdin.read()
        saved = extract_and_save_memories(text)
        print(json.dumps({"saved": saved}))
        return

    if args.command == "clear-memory":
        clear_memory()
        print(json.dumps({"ok": True}) if args.json else "Memory cleared.")
        return

    if args.command == "clear-knowledge":
        clear_knowledge()
        print(json.dumps({"ok": True}) if args.json else "Knowledge base cleared.")
        return


if __name__ == "__main__":
    main()
